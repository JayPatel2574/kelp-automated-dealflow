from pptx.util import Inches
from src.ppt.theme import FONT_BODY, BODY_SIZE
from src.ppt.colors import KELP_DARK_GREY
from src.ppt.branding import add_logo_text, add_footer
from src.ppt.layout_engine import white_content_panel


def build_overview(prs, slide, data, overview_text):
    white_content_panel(slide, Inches(0.5), Inches(1.2), Inches(9), Inches(4.8))

    box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.5),
        Inches(8.4), Inches(4)
    )
    tf = box.text_frame
    tf.word_wrap = True

    bullets = [
        overview_text,
        f"Website: {data['company_profile']['website']}"
    ]

    for i, text in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = text
        p.font.name = FONT_BODY
        p.font.size = BODY_SIZE
        p.font.color.rgb = KELP_DARK_GREY

    add_logo_text(slide)
    add_footer(slide)
