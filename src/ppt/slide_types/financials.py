from pptx.util import Inches
from src.ppt.colors import KELP_DARK_GREY
from src.ppt.branding import add_logo_text, add_footer
from src.ppt.theme import FONT_BODY, BODY_SIZE
from src.ppt.layout_engine import white_content_panel


def build_financials(prs, slide, data):
    white_content_panel(slide, Inches(0.5), Inches(1.2), Inches(9), Inches(4.8))

    income = data["financials"].get("Income Statement", {})
    revenue = income.get("Revenue From Operations", {})
    pat = income.get("PAT", {})

    latest_year = sorted(revenue.keys())[-1]
    revenue_latest = revenue.get(latest_year)
    pat_latest = pat.get(latest_year)

    metrics = [
        f"Revenue ({latest_year}): ₹{revenue_latest:.1f} Cr",
        f"PAT ({latest_year}): ₹{pat_latest:.1f} Cr"
    ]

    box = slide.shapes.add_textbox(
        Inches(1), Inches(1.8),
        Inches(8), Inches(3)
    )
    tf = box.text_frame

    for i, metric in enumerate(metrics):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = metric
        p.font.name = FONT_BODY
        p.font.size = BODY_SIZE
        p.font.color.rgb = KELP_DARK_GREY

    add_logo_text(slide)
    add_footer(slide)
