from pptx.util import Inches
from src.ppt.colors import KELP_WHITE, KELP_DARK_INDIGO
from src.ppt.branding import add_logo_text, add_footer
from src.ppt.theme import FONT_BODY, BODY_SIZE
from src.ppt.layout_engine import full_bleed_background


def build_highlights(prs, slide, insights):
    full_bleed_background(slide, prs, KELP_DARK_INDIGO)

    box = slide.shapes.add_textbox(
        Inches(0.8), Inches(1.5),
        Inches(8.4), Inches(4.5)
    )
    tf = box.text_frame

    for i, highlight in enumerate(insights["highlights"][:4]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {highlight}"
        p.font.name = FONT_BODY
        p.font.size = BODY_SIZE
        p.font.color.rgb = KELP_WHITE

    add_logo_text(slide)
    add_footer(slide)
