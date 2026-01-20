from pptx.util import Inches
from src.ppt.colors import KELP_DARK_GREY
from src.ppt.branding import add_logo_text, add_footer
from src.ppt.theme import FONT_BODY, BODY_SIZE


DISCLAIMER_TEXT = (
    "This presentation is strictly private and confidential and has been prepared by "
    "Kelp M&A Team for discussion purposes only. It does not constitute an offer or "
    "solicitation. All information is based on sources believed to be reliable."
)


def build_disclaimer(prs, slide):
    box = slide.shapes.add_textbox(
        Inches(1), Inches(2),
        Inches(8), Inches(3)
    )
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = DISCLAIMER_TEXT
    p.font.name = FONT_BODY
    p.font.size = BODY_SIZE
    p.font.color.rgb = KELP_DARK_GREY

    add_logo_text(slide)
    add_footer(slide)
